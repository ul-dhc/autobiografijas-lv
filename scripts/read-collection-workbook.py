"""Read selected catalogue fields only. Never copy workbook notes or private fields."""
import json
import re
import sys
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1])
DATE = '2026-10-14'
DIR = ROOT / 'docs' / 'collections'
index = json.loads((DIR / f'garamantas-list-lv-{DATE}.json').read_text())['rows']
web = {r['cells'][1]['text'].strip(): r for r in index}
workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook['Kolekcijas']
rows = defaultdict(list)
for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
    if not row or not isinstance(row[0], (int, float)) or int(row[0]) != row[0]:
        continue
    # Restrict retained cells to number, title, submitter, writing dates, material type.
    rows[str(int(row[0]))].append({
        'row': row_number, 'title': str(row[1] or '').strip(),
        'submitter': str(row[3] or '').strip(),
        'start': row[5], 'end': row[6], 'type': row[11],
    })

def normalize(value):
    return re.sub(r'[^\w]+', ' ', str(value or '').casefold()).strip()

def year(value):
    text = str(value or '').strip()
    if not re.fullmatch(r'\d{4}(?:\.0)?', text):
        return None
    result = int(float(text))
    return result if 1000 <= result <= 2026 else None

type_map = {'D': 'diary', 'A': 'autobiography', 'V': 'letters'}
supplement = {}
audit = {'sourceSheet': 'Kolekcijas', 'sourceRows': sum(map(len, rows.values())),
         'uniqueNumbers': len(rows), 'unmatchedNumbers': sorted(set(rows) - set(web), key=int),
         'duplicateNumbers': [number for number, rs in rows.items() if len(rs) > 1],
         'skippedTitleNumbers': [], 'ambiguousNumbers': [], 'reversedDateNumbers': [],
         'submitterRowsWithoutCorroboration': [], 'matchedRows': {}}
for number, public in web.items():
    candidates = rows.get(number, [])
    if not candidates:
        continue
    title = normalize(public['cells'][6]['text'])
    exact = [r for r in candidates if normalize(r['title']) == title]
    if len(exact) == 1:
        chosen = exact[0]
    elif len(candidates) == 1:
        chosen = candidates[0]
        if SequenceMatcher(None, title, normalize(chosen['title'])).ratio() < .65:
            audit['skippedTitleNumbers'].append(number)
            continue
    else:
        audit['ambiguousNumbers'].append(number)
        continue
    start, end = year(chosen['start']), year(chosen['end'])
    if start and end and start > end:
        audit['reversedDateNumbers'].append(number)
        start = end = None
    parts = re.split(r'\s*(?:,|;|\n|\s+un\s+)\s*', chosen['submitter'])
    confirmed = []
    for link in public['cells'][8]['links']:
        name = link['text'].strip()
        if name and any(normalize(name) == normalize(part) for part in parts):
            # Names only: no contact information or relationship notes.
            confirmed.append(name)
    if chosen['submitter'] and not confirmed:
        audit['submitterRowsWithoutCorroboration'].append(number)
    codes = str(chosen['type'] or '').split('/')
    types = [type_map[code] for code in codes if code in type_map]
    supplement[number] = {
        'writingPeriod': {'start': start, 'end': end},
        'materialTypes': types,
        'confirmedSubmitters': sorted(set(confirmed)),
    }
    audit['matchedRows'][number] = chosen['row']

audit['matchedCollections'] = len(supplement)
audit['collectionsWithConfirmedSubmitters'] = sum(bool(r['confirmedSubmitters']) for r in supplement.values())
audit['collectionsWithWritingPeriod'] = sum(any(r['writingPeriod'].values()) for r in supplement.values())
audit['excludedColumns'] = ['Autors', 'Iesniegšanas datums', 'Aprakstītā sākums', 'Aprakstītā beigas',
                           'Dzimums', 'Kolektīvā/Individuālā/vēstules', 'Tehniska/saturiska informācija',
                           'Vienošanās', 'Oriģināls/kopija', 'Statuss', 'TEKSTS',
                           'Teksta sarežģītības pakāpe (1-5)', 'Komentāri']
audit['sourceFile'] = SOURCE.name
audit['retrievedOn'] = DATE
(DIR / 'workbook-public-supplement.json').write_text(json.dumps(supplement, ensure_ascii=False, indent=2) + '\n')
(DIR / 'workbook-import-audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({key: audit[key] for key in ('matchedCollections', 'collectionsWithConfirmedSubmitters',
      'collectionsWithWritingPeriod', 'skippedTitleNumbers', 'ambiguousNumbers', 'reversedDateNumbers')}, indent=2))
